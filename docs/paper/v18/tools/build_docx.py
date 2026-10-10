# -*- coding: utf-8 -*-
"""Build the v18 manuscript as an IEEE-style two-column .docx.

    python build_docx.py <figures_dir> <output.docx>

Convert to PDF with:  soffice --headless --convert-to pdf <output.docx>
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
import manuscript as M  # noqa: E402

FIG = Path(sys.argv[1])
OUT = Path(sys.argv[2])
FONT = "Times New Roman"
DEVA = "Noto Serif Devanagari"

ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]

# ── citation numbering by first appearance ─────────────────────────────────
order: list[str] = []


def _scan(text: str) -> None:
    groups = re.findall(r"\[@[A-Za-z0-9]+(?:\s*,\s*@[A-Za-z0-9]+)*\]", text)
    for key in re.findall(r"@([A-Za-z0-9]+)", " ".join(groups)):
        if key not in order:
            if key not in M.REFS:
                raise KeyError(f"unknown reference key {key}")
            order.append(key)


def _scan_block(kind, data):
    if kind in ("p", "ack"):
        _scan(data)
    elif kind == "bullets":
        for b in data:
            _scan(b)
    elif kind == "table":
        for row in data["rows"]:
            for c in row:
                _scan(c)
        _scan(data["caption"])


for kind, data in M.BODY:
    _scan_block(kind, data)
_scan_block("table", M.TA1)
missing = [k for k in M.REFS if k not in order]
if missing:
    raise SystemExit(f"references never cited: {missing}")


def cite(text: str) -> str:
    def one(m):
        keys = re.findall(r"@([A-Za-z0-9]+)", m.group(0))
        return ", ".join(f"[{order.index(k) + 1}]" for k in keys)
    return re.sub(r"\[@[A-Za-z0-9]+(?:\s*,\s*@[A-Za-z0-9]+)*\]", one, text)


# ── table numbering (IEEE Roman numerals) ──────────────────────────────────
table_ids = []
for kind, data in M.BODY:
    if kind == "table":
        table_ids.append(re.match(r"Table (\d+)\.", data["caption"]).group(1))


def roman_tables(text: str) -> str:
    return re.sub(r"\bTables? (\d+)(?=\b)",
                  lambda m: m.group(0).replace(m.group(1), ROMAN[int(m.group(1))]), text)


def prep(text: str) -> str:
    text = cite(text)
    text = re.sub(r"Tables (\d+) and (\d+)",
                  lambda m: f"Tables {ROMAN[int(m.group(1))]} and {ROMAN[int(m.group(2))]}", text)
    text = re.sub(r"Table (\d+)\b", lambda m: f"Table {ROMAN[int(m.group(1))]}", text)
    return text


# ── low-level formatting helpers ───────────────────────────────────────────
def set_font(run, size, bold=None, italic=None, mono=False):
    run.font.name = "DejaVu Sans Mono" if mono else FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    base = "DejaVu Sans Mono" if mono else FONT
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), base)
    rfonts.set(qn("w:cs"), DEVA)
    run.font.size = Pt(size)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(szcs)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_rich(par, text, size=10, bold=False, italic=False):
    for piece in TOKEN.split(text):
        if not piece:
            continue
        b, i, mono = bold, italic, False
        if piece.startswith("**") and piece.endswith("**"):
            piece, b = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 1:
            piece, i = piece[1:-1], not italic
        elif piece.startswith("`") and piece.endswith("`"):
            piece, mono = piece[1:-1], True
        run = par.add_run(piece)
        set_font(run, size - (1 if mono else 0), bold=b, italic=i, mono=mono)
    return par


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, indent=None,
        left=None, hanging=None, keep=False):
    pf = par.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.widow_control = True
    if indent is not None:
        pf.first_line_indent = Inches(indent)
    if left is not None:
        pf.left_indent = Inches(left)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    if keep:
        pf.keep_with_next = True
    return par


def set_columns(section, num, space=0.25):
    sectpr = section._sectPr
    cols = sectpr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectpr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space * 1440)))


def page_setup(section):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(1.0)
    section.left_margin = section.right_margin = Inches(0.625)


def new_section(doc, cols):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    page_setup(sec)
    set_columns(sec, cols)
    return sec


def cell_border(cell, **kw):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge, val in kw.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val.get("val", "single"))
        el.set(qn("w:sz"), str(val.get("sz", 6)))
        el.set(qn("w:color"), val.get("color", "000000"))
        borders.append(el)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hex_fill)
    tcpr.append(sh)


def cell_margins(table, top=20, bottom=20, left=50, right=50):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def fix_grid(table, widths):
    """LibreOffice sizes columns from tblGrid and tblW, not from cell widths."""
    tbl = table._tbl
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)


# ── block renderers ────────────────────────────────────────────────────────
h1_count = 0
h2_count = 0


def h1(doc, text):
    global h1_count, h2_count
    h1_count += 1
    h2_count = 0
    par = doc.add_paragraph()
    fmt(par, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
    run = par.add_run(f"{ROMAN[h1_count]}. ")
    set_font(run, 10)
    run = par.add_run(text)
    set_font(run, 10)
    run.font.small_caps = True


def h2(doc, text):
    global h2_count
    h2_count += 1
    par = doc.add_paragraph()
    fmt(par, WD_ALIGN_PARAGRAPH.LEFT, before=5, after=2, keep=True)
    run = par.add_run(f"{chr(64 + h2_count)}. {text}")
    set_font(run, 10, italic=True)


def para(doc, text, size=10, indent=0.17):
    par = doc.add_paragraph()
    fmt(par, indent=indent)
    add_rich(par, prep(text), size)
    return par


def equation(doc, expr, num):
    par = doc.add_paragraph()
    fmt(par, WD_ALIGN_PARAGRAPH.LEFT, before=3, after=3, left=0.12)
    pf = par.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(3.38), alignment=2)  # right
    lines = expr.split("\n")
    for i, line in enumerate(lines):
        run = par.add_run(line)
        set_font(run, 9, italic=True)
        if i < len(lines) - 1:
            run.add_break()
    run = par.add_run(f"\t({num})")
    set_font(run, 9)


def algorithm(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Inches(3.35)
    fix_grid(t, [3.35])
    shade(cell, "F5F5F5")
    b = {"val": "single", "sz": 6}
    cell_border(cell, top=b, bottom=b, left=b, right=b)
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        fmt(par, WD_ALIGN_PARAGRAPH.LEFT, after=(2 if i == 0 else 0))
        add_rich(par, line, 8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, spec):
    par = doc.add_paragraph()
    fmt(par, WD_ALIGN_PARAGRAPH.CENTER, before=4, after=2, keep=True)
    par.add_run().add_picture(str(FIG / spec["file"]), width=Inches(spec["width"]))
    cap = doc.add_paragraph()
    fmt(cap, WD_ALIGN_PARAGRAPH.JUSTIFY, after=6)
    add_rich(cap, prep(spec["caption"]), 8)


def table(doc, spec, label=None):
    wide = spec.get("wide")
    if wide:
        new_section(doc, 1)
    size = spec.get("font", 7.5)
    m = re.match(r"Table (A?\d+)\.\s*(.*)", spec["caption"], re.S)
    num, rest = m.group(1), m.group(2)
    lab = num if num.startswith("A") else ROMAN[int(num)]
    cap = doc.add_paragraph()
    fmt(cap, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=1, keep=True)
    run = cap.add_run(f"TABLE {lab}")
    set_font(run, 8)
    cap2 = doc.add_paragraph()
    fmt(cap2, WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
    add_rich(cap2, prep(rest), 8)
    cap2.runs[0].font.small_caps = False

    ncol = len(spec["header"])
    t = doc.add_table(rows=1 + len(spec["rows"]), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t)
    fix_grid(t, spec["widths"])
    thick, thin = {"sz": 10}, {"sz": 4}
    for r, row in enumerate([spec["header"]] + spec["rows"]):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Inches(spec["widths"][c])
            par = cell.paragraphs[0]
            numeric = c > 0 and re.fullmatch(r"[\d.,%+−— ]+", val or "") is not None
            fmt(par, WD_ALIGN_PARAGRAPH.CENTER if numeric else WD_ALIGN_PARAGRAPH.LEFT)
            add_rich(par, prep(val), size, bold=(r == 0))
            edges = {}
            if r == 0:
                edges.update(top=thick, bottom=thin)
                shade(cell, "EDEDED")
            if r == len(spec["rows"]):
                edges["bottom"] = thick
            if edges:
                cell_border(cell, **edges)
            if r < len(spec["rows"]):
                par.paragraph_format.keep_with_next = True
        # keep rows together where possible
        trpr = t.rows[r]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
    if spec.get("note"):
        np_ = doc.add_paragraph()
        fmt(np_, WD_ALIGN_PARAGRAPH.LEFT, before=2)
        add_rich(np_, prep(spec["note"]), 7.5, italic=True)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    if wide:
        new_section(doc, 2)


def bullets(doc, items):
    for i, text in enumerate(items, 1):
        par = doc.add_paragraph()
        fmt(par, left=0.22, hanging=0.2, after=1)
        run = par.add_run(f"{i}) ")
        set_font(run, 10)
        add_rich(par, prep(text), 10)


def references(doc):
    par = doc.add_paragraph()
    fmt(par, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
    run = par.add_run("References")
    set_font(run, 10)
    run.font.small_caps = True
    for i, key in enumerate(order, 1):
        par = doc.add_paragraph()
        fmt(par, WD_ALIGN_PARAGRAPH.JUSTIFY, left=0.3, hanging=0.3, after=1)
        run = par.add_run(f"[{i}]\t")
        set_font(run, 8)
        par.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
        add_rich(par, M.REFS[key], 8)


# ── document ───────────────────────────────────────────────────────────────
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
page_setup(doc.sections[0])
set_columns(doc.sections[0], 1)

p = doc.add_paragraph()
fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=10)
set_font(p.add_run(M.TITLE), 24)
p = doc.add_paragraph()
fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=3)
set_font(p.add_run(M.AUTHORS), 11)
p = doc.add_paragraph()
fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=2)
set_font(p.add_run(M.AFFIL), 10, italic=True)
p = doc.add_paragraph()
fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=10)
set_font(p.add_run(M.EMAILS), 9)

new_section(doc, 2)
p = doc.add_paragraph()
fmt(p, after=4)
set_font(p.add_run("Abstract"), 9, bold=True, italic=True)
set_font(p.add_run("—" + M.ABSTRACT), 9, bold=True)
p = doc.add_paragraph()
fmt(p, after=6)
set_font(p.add_run("Index Terms"), 9, bold=True, italic=True)
set_font(p.add_run("—" + M.KEYWORDS), 9, bold=True)

for kind, data in M.BODY:
    if kind == "h1":
        h1(doc, data)
    elif kind == "h2":
        h2(doc, data)
    elif kind == "p":
        para(doc, data)
    elif kind == "eq":
        equation(doc, *data)
    elif kind == "algo":
        algorithm(doc, data)
    elif kind == "fig":
        figure(doc, data)
    elif kind == "table":
        table(doc, data)
    elif kind == "bullets":
        bullets(doc, data)
    elif kind == "ack":
        par = doc.add_paragraph()
        fmt(par, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
        r = par.add_run("Acknowledgment")
        set_font(r, 10)
        r.font.small_caps = True
        para(doc, data)

references(doc)

# Appendix (single column)
new_section(doc, 1)
par = doc.add_paragraph()
fmt(par, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
r = par.add_run("Appendix: Scheme Validation Status")
set_font(r, 10)
r.font.small_caps = True
table(doc, dict(M.TA1, wide=False))

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("saved", OUT, "with", len(order), "references")
